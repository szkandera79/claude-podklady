import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
db=json.load(open('svj_db.json')); hits=json.load(open('street_hits.json'))
KR={}
def keep(s):
    n=s['obchodniJmeno']; nl=n.lower()
    if 'likvidac' in nl: return False
    reg=s.get('seznamRegistraci',{})
    if reg.get('stavZdrojeRes') not in (None,'AKTIVNI'): return False
    if s['pravniForma']=='145': return True
    nace=set(s.get('czNace',[])+s.get('czNace2008',[]))
    return 'bytov' in nl or bool(nace & {'6820','68200','6832','68320'})
def typ(s): return 'SVJ' if s['pravniForma']=='145' else 'Bytové družstvo'
def adr(s):
    a=s['sidlo']; ul=a.get('nazevUlice') or a.get('nazevCastiObce') or ''
    cp=str(a.get('cisloDomovni') or ''); co=str(a.get('cisloOrientacni') or '')+(a.get('cisloOrientacniPismeno') or '')
    return (f"{ul} {cp}"+(f"/{co}" if co else '')).strip()
def row(s,cena,zdroj,shoda,prio):
    a=s['sidlo']
    return [prio,s['obchodniJmeno'],s['ico'],typ(s),adr(s),a.get('nazevCastiObce') or '',a.get('nazevObce') or '',
            a.get('psc') or '',a.get('nazevKraje') or '',round(cena),round(cena/1.12),zdroj,shoda,s.get('datumVzniku',''),
            f"https://ares.gov.cz/ekonomicke-subjekty?ico={s['ico']}",
            f"https://or.justice.cz/ias/ui/rejstrik-$firma?ico={s['ico']}"]
H=['Priorita','Název','IČO','Typ','Adresa sídla (dům)','Část obce','Obec','PSČ','Kraj','Cena tepla CZT Kč/GJ s DPH','bez DPH','Zdroj ceny (lokalita ERÚ 2026)','Typ shody','Vznik','ARES','Veřejný rejstřík']
A=[];seen=set()
for h in hits:
    s=h['subj']
    if not keep(s) or s['ico'] in seen: continue
    seen.add(s['ico']); A.append(row(s,h['cena'],h['lok'],'Adresa domu je přímo v ceníku ERÚ','A'))
A.sort(key=lambda r:-r[9])
B=[];summ=[]
for obec,v in db.items():
    cena,zdroj=v['info']
    subs=[s for s in v['subs'] if keep(s) and s['ico'] not in seen]
    prio='B' if cena>=1250 else 'C'
    for s in subs: B.append(row(s,cena,zdroj,'Obec s drahým CZT — napojení domu ověřit',prio))
    summ.append([obec,cena,round(cena/1.12),sum(typ(s)=='SVJ' for s in subs),sum(typ(s)!='SVJ' for s in subs),len(subs),zdroj])
B.sort(key=lambda r:(-r[9],r[6],r[4]))
summ.sort(key=lambda r:-r[1])
wb=Workbook()
def sheet(ws,head,rows,widths):
    ws.append(head)
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='0F3D2E'); c.alignment=Alignment(wrap_text=True,vertical='top')
    for r in rows: ws.append(r)
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
W=[8,55,11,15,32,20,16,8,20,12,10,40,34,11,18,18]
ws=wb.active; ws.title='A – přesné adresy'; sheet(ws,H,A,W)
ws=wb.create_sheet('B+C – obce s drahým CZT'); sheet(ws,H,B,W)
ws=wb.create_sheet('Souhrn podle obcí'); sheet(ws,['Obec','Cena CZT s DPH','bez DPH','SVJ','Bytová družstva','Celkem','Zdroj ceny (ERÚ)'],summ,[24,14,10,8,14,9,50])
ws=wb.create_sheet('Jak s tím pracovat')
notes=[
 "Seznam společenství vlastníků jednotek (SVJ) a bytových družstev v lokalitách s drahým teplem z CZT.",
 "Zdroje: ceny tepla – ERÚ, předběžné ceny k 1. 1. 2026 (domovní předávací stanice / sekundární rozvod, s DPH); subjekty – ARES (veřejný rejstřík), stav k 22. 9. 2026.",
 "Priorita A: adresa domu je přímo uvedena jako cenová lokalita v ceníku ERÚ – cena tepla se týká přímo tohoto domu (některé jsou domovní kotelny dodavatele).",
 "Priorita B: obec s cenou CZT ≥ 1 250 Kč/GJ s DPH. Priorita C: 1 140–1 250 Kč/GJ. U B a C není jisté, že je dům na CZT napojen – ověřit (mapa soustavy dodavatele, prohlídka, dotaz na výbor).",
 "Družstva jsou filtrována na bytová (název nebo činnost správa/pronájem nemovitostí); SVJ a družstva v likvidaci vyřazena.",
 "",
 "KONTAKT – e‑maily ani datové schránky v seznamu záměrně nejsou:",
 "• SVJ e‑mail veřejně neuvádějí a ARES ho neobsahuje.",
 "• Rozesílání nevyžádaných obchodních sdělení e‑mailem bez předchozího souhlasu zakazuje § 7 zákona č. 480/2004 Sb. – i vůči právnickým osobám (dozor ÚOOÚ, pokuty).",
 "• Obchodní sdělení do datové schránky bez souhlasu je rovněž přestupek (řeší Digitální a informační agentura).",
 "• Bezpečná cesta: klasický dopis na adresu sídla (= adresa domu) k rukám výboru SVJ / představenstva družstva, s vzorovou nabídkou a výzvou ‚pošlete nám vyúčtování‘. Souhlas s e‑mailovou komunikací pak získáte od těch, kdo se ozvou.",
 "• Další cesta: správcovské firmy, které SVJ spravují – partnerská nabídka firmě (i zde se vyhnout hromadnému nevyžádanému e‑mailu).",
]
for n in notes: ws.append([n])
ws.column_dimensions['A'].width=160
out='../SVJ a BD – drahé teplo CZT 2026.xlsx'
wb.save(out)
print(len(A),len(B)); print(sum(r[0]=='B' for r in B), sum(r[0]=='C' for r in B))
for s in summ: print(s[:6])
