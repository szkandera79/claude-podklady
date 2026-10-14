#!/usr/bin/env python3
"""Podklady pro registraci (SVT / Nová zelená úsporám, EPREL) – tepelná čerpadla řady REO R290 Smart, 13 modelů.

Dvě varianty:
  A_REO-HEATING/  – značka REO HEATING, označení RH06 1F … RH16 3F, RHX20 Pro, RHX30 Pro; Ria Trade jako dodavatel
  B_Power-World/  – značka a označení výrobce (PWxx-DKXLRE1/3-H8), výrobce Power World, Ria Trade jako dovozce;
                    navazuje na certifikáty HP KEYMARK výrobce (BRE Global, 041-K032-23…31)

Každá varianta: formuláře dle nařízení (EU) 811/2013 (xlsx + pdf, listy 35 °C a 55 °C), energetické štítky (pdf),
přehled modelů (xlsx); varianta B navíc certifikáty KEYMARK.

    python3 build.py            # vše kromě PDF formulářů
    python3 build.py --pdf      # + export formulářů do PDF přes Microsoft Excel
"""
import glob
import html
import os
import re
import shutil
import subprocess
import sys
import zipfile

import fitz
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

DIR = os.path.dirname(os.path.abspath(__file__))
PWH = os.path.join(DIR, '..', '..')                               # REO/PW Helios
SRC_FORMS = os.path.join(PWH, 'EU regulation SCOP Helios')
SRC_LABELS = os.path.join(PWH, 'Helios Energy Label and ERP report')
SRC_CERTS = os.path.join(PWH, '..', 'R290 new models', 'H8 Keymark')
A, B = os.path.join(DIR, 'A_REO-HEATING'), os.path.join(DIR, 'B_Power-World')

RIA = 'Ria Trade, s.r.o., Korunní 2569/108, 101 00 Praha 10, Czech Republic, www.reoheating.cz'
PW_ADDR = ('Power World Machinery Equipment Co. Ltd., No.24, the Fourth Industrial Zone, HouTing Street, '
           'ShaJing Town, BaoAn District, Shenzhen, P.R. China')
CONTACT = {'A': RIA, 'B': f'Manufacturer: {PW_ADDR}. Importer / supplier: {RIA}'}
ADDR_OLD = re.compile(r'Ria Trade, s\.r\.o\. Korunn(?:í|&#237;) 2569/108, 101 00, Czech Republic www\.reoheating\.cz')
CERT = {'PW06-DKXLRE1-H8': '041-K032-23', 'PW08-DKXLRE1-H8': '041-K032-24', 'PW08-DKXLRE3-H8': '041-K032-24',
        'PW10-DKXLRE1-H8': '041-K032-25', 'PW10-DKXLRE3-H8': '041-K032-26', 'PW12-DKXLRE1-H8': '041-K032-27',
        'PW12-DKXLRE3-H8': '041-K032-27', 'PW14-DKXLRE1-H8': '041-K032-28', 'PW14-DKXLRE3-H8': '041-K032-28',
        'PW16-DKXLRE1-H8': '041-K032-29', 'PW16-DKXLRE3-H8': '041-K032-29', 'PW20-DKXLRE3-H8': '041-K032-30',
        'PW30-DKXLRE3-H8': '041-K032-31'}
BAN_A = re.compile(r'(?i)power\s*world|helios|PW\d\d-|DKXL|派沃|Beryl|KSOProduct|Feli.ski')


def reo(code):
    n, ph = re.match(r'PW(\d\d)-DKXLRE(\d)-H8', code).groups()
    return {'20': 'RHX20 Pro', '30': 'RHX30 Pro'}.get(n, f'RH{n} {ph}F')


def fname(model):
    return model.replace(' ', '-')


CORE = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<cp:coreProperties '
        'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        '<dc:title>{t}</dc:title><dc:creator>{a}</dc:creator><cp:lastModifiedBy>{a}</cp:lastModifiedBy>'
        '<dcterms:created xsi:type="dcterms:W3CDTF">2026-10-14T00:00:00Z</dcterms:created>'
        '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-10-14T00:00:00Z</dcterms:modified></cp:coreProperties>')


# ---------------------------------------------------------------- formuláře 811/2013
def form(src, var, code):
    model = reo(code) if var == 'A' else code
    root = A if var == 'A' else B
    out = os.path.join(root, 'formulare-811-2013', f'{fname(model)}_EU-811-2013_parametry-35-55C.xlsx')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    zin = zipfile.ZipFile(src)
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == 'docProps/custom.xml':
                continue                                            # vlastnosti WPS (KSOProductBuildVer, ICV)
            if it.filename.endswith(('.xml', '.rels')):
                s = data.decode('utf8')
                if it.filename == 'xl/sharedStrings.xml' or it.filename.startswith('xl/worksheets/'):
                    # každý kód modelu ve formuláři → model ze jména souboru
                    # (opraví i chybu výrobce: list 55 °C ve formuláři 30 kW nesl označení PW20)
                    s = re.sub(r'PW\d\d-DKXLRE\d-H8', model, s)
                    s = ADDR_OLD.sub(html.escape(CONTACT[var], quote=False), s)
                if it.filename.startswith('xl/worksheets/sheet'):
                    # tisk každého listu na jednu stranu A4 (originál přetékal na 2. stranu)
                    s = s.replace('<sheetPr/>', '<sheetPr><pageSetUpPr fitToPage="1"/></sheetPr>')
                    s = s.replace('<pageSetUpPr/>', '<pageSetUpPr fitToPage="1"/>')
                    s = re.sub(r'<pageSetup ', '<pageSetup fitToWidth="1" fitToHeight="1" ', s, count=1)
                    if var == 'B':                                  # delší kontakt (výrobce + dovozce) → vyšší řádek
                        s = re.sub(r'(<row r="38"[^>]*?)ht="27"', r'\1ht="46"', s)
                if it.filename == 'docProps/core.xml':
                    s = CORE.format(t=f'{model} – technical parameters, Regulation (EU) 811/2013',
                                    a='REO HEATING' if var == 'A' else 'Ria Trade, s.r.o.')
                if it.filename == 'docProps/app.xml':
                    s = s.replace('工作表', 'Worksheets')
                if it.filename == '[Content_Types].xml':
                    s = re.sub(r'<Override PartName="/docProps/custom.xml"[^>]*/>', '', s)
                if it.filename == '_rels/.rels':
                    s = re.sub(r'<Relationship [^>]*Target="docProps/custom.xml"[^>]*/>', '', s)
                data = s.encode('utf8')
            zout.writestr(it, data)
    z = zipfile.ZipFile(out)
    txt = html.unescape(''.join(z.read(n).decode('utf8', 'ignore') for n in z.namelist() if n.endswith(('.xml', '.rels'))))
    if var == 'A':
        assert not BAN_A.findall(txt), (out, set(BAN_A.findall(txt)))
    else:
        assert 'PW20-DKXLRE3-H8' not in txt or code == 'PW20-DKXLRE3-H8', out
    assert model in txt and CONTACT[var] in txt, out
    return out


# ---------------------------------------------------------------- energetické štítky
def label_src(code):
    n, ph = re.match(r'PW(\d\d)-DKXLRE(\d)-H8', code).groups()
    for f in glob.glob(os.path.join(SRC_LABELS, f'PW{n}*.pdf')):
        b = os.path.basename(f).lower()
        if n in ('20', '30') or (ph == '1' and 'single' in b) or (ph == '3' and '3' in b.replace(f'pw{n}', '')):
            return f
    raise SystemExit(f'štítek pro {code} nenalezen')


def label(code, var):
    src = label_src(code)
    model = reo(code) if var == 'A' else code
    out = os.path.join(A if var == 'A' else B, 'energeticke-stitky', f'{fname(model)}_energeticky-stitek.pdf')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    if var == 'B':
        shutil.copyfile(src, out)                                   # originál výrobce beze změny
        return out
    # Varianta A: přepis textových operátorů v obsahu stránky (štítek je regulovaný formát – grafika, třídy
    # a čísla se nemění). „POWER WORLD“ → „REO HEATING“; model (v originále skládaný z několika fontů)
    # → jeden řetězec ve fontu použitém pro značku (vložený celý Calibri), stejný začátek řádku.
    doc = fitz.open(src)
    page = doc[0]
    done = False
    for x in page.get_contents():
        s = doc.xref_stream(x).decode('latin-1')
        m = re.search(r'\[\(POWER WORLD\)\s*(-?[\d.]+)\s*\(PW\)\]TJ', s)
        if m:                                                       # značka a začátek modelu v jednom TJ s mezerou
            calibri = next(fitz.Font(fontbuffer=doc.extract_font(f[0])[3]) for f in page.get_fonts()
                           if f[3] == 'Calibri' and doc.extract_font(f[0])[3])
            dx = calibri.text_length('POWER WORLD', 11) - float(m.group(1)) * 11 / 1000
            j = s.index('ET', m.end())
            s = s[:m.start()] + f'(REO HEATING)Tj\n{dx:.3f} 0 Td\n({model})Tj\n' + s[j:]
        elif '(POWER WORLD)Tj' in s:
            assert s.count('(POWER WORLD)Tj') == 1 and s.count('(PW)Tj') == 1, src
            s = s.replace('(POWER WORLD)Tj', '(REO HEATING)Tj')
            i = s.index('(PW)Tj')
            j = s.index('ET', i)
            s = s[:i] + f'({model})Tj\n' + s[j:]
        else:
            continue
        doc.update_stream(x, s.encode('latin-1'))
        done = True
    assert done, src
    doc.set_metadata({'title': f'{model} – energetický štítek', 'author': 'REO HEATING',
                      'subject': 'Energetický štítek podle nařízení (EU) 811/2013',
                      'keywords': f'{model}, REO HEATING, REO R290 Smart', 'creator': 'REO HEATING',
                      'producer': 'REO HEATING'})
    doc.save(out, garbage=3, deflate=True)
    t = fitz.open(out)[0].get_text()
    assert 'REO HEATING' in t and model in t and not BAN_A.findall(t), (out, t[:200])
    orig = fitz.open(src)[0].get_text()
    for v in re.findall(r'\b\d{2}\b', orig):                        # hodnoty výkonu a hluku zůstaly
        assert v in t, (out, v)
    return out


# ---------------------------------------------------------------- přehled modelů
def values(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    r = {}
    for ws in wb:
        t = '35' if '35' in ws.title else '55'
        for row in ws.iter_rows(values_only=True):
            v = [x for x in row if x not in (None, '')]
            if not v:
                continue
            k = str(v[0])
            if k.startswith('Rated heat output'):
                r['Prated'] = v[2]
                r['ns' + t] = v[6]
            if k.startswith('Sound power level'):
                r['LWA' + t] = str(v[2]).split('/')[-1]
            if k.startswith('Annual energy consumption'):
                r['QHE' + t] = v[2]
    return r


def summary(var, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Přehled modelů'
    hdr = ['Model', 'Napájení', 'Jmenovitý výkon Prated (kW)', 'ηs 35 °C (%)', 'SCOP 35 °C', 'Třída 35 °C',
           'ηs 55 °C (%)', 'SCOP 55 °C', 'Třída 55 °C', 'Hlučnost LWA 35/55 (dB)', 'Roční spotřeba 35 °C (kWh)',
           'Formulář 811/2013', 'Energetický štítek']
    if var == 'B':
        hdr[0] = 'Model (výrobce)'
        hdr += ['Certifikát HP KEYMARK', 'Držitel certifikátu']
    ws.append(hdr)
    cls = lambda ns, t: ('A+++' if ns >= (175 if t == 35 else 150) else 'A++' if ns >= (150 if t == 35 else 125) else 'A+')
    for code, form_path, label_path in rows:
        v = values(form_path)
        model = reo(code) if var == 'A' else code
        line = [model, '1 fáze 230 V' if 'E1' in code else '3 fáze 400 V', v['Prated'],
                v['ns35'], round((v['ns35'] + 3) * 2.5 / 100, 2), cls(v['ns35'], 35),
                v['ns55'], round((v['ns55'] + 3) * 2.5 / 100, 2), cls(v['ns55'], 55),
                f"{v['LWA35']} / {v['LWA55']}", v['QHE35'],
                os.path.basename(form_path), os.path.basename(label_path)]
        if var == 'B':
            line += [CERT[code], 'Power World Machinery Equipment Co. Ltd. (BRE Global, platnost do 13. 8. 2036)']
        ws.append(line)
    green = PatternFill('solid', fgColor='34A263')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = green
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for col, w in zip('ABCDEFGHIJKLMNO', [16, 13, 12, 10, 9, 9, 10, 9, 9, 12, 13, 42, 38, 18, 40]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 45
    note = ('SCOP = (ηs + 3 %) × 2,5 (průměrné klima, EN 14825); třída dle nařízení (EU) 811/2013. '
            'Dodavatel: ' + RIA + '.')
    if var == 'B':
        note += ' Výrobce: ' + PW_ADDR + '. Ria Trade není výrobcem – uvádí výrobky na trh jako dovozce.'
    ws.append([])
    ws.append([note])
    out = os.path.join(A if var == 'A' else B, 'prehled-modelu.xlsx')
    wb.properties.creator = 'REO HEATING' if var == 'A' else 'Ria Trade, s.r.o.'
    wb.save(out)
    return out


# ---------------------------------------------------------------- export formulářů do PDF (Excel)
def export_pdf(root):
    sandbox = os.path.expanduser('~/Library/Containers/com.microsoft.Excel/Data/Documents/reo-export')
    os.makedirs(sandbox, exist_ok=True)
    for f in sorted(glob.glob(os.path.join(root, 'formulare-811-2013', '*.xlsx'))):
        name = os.path.basename(f)
        src = os.path.join(sandbox, name)
        pdf = src[:-5] + '.pdf'
        shutil.copyfile(f, src)
        if os.path.exists(pdf):
            os.remove(pdf)
        script = f'''with timeout of 300 seconds
tell application "Microsoft Excel"
  set wb to open workbook workbook file name (POSIX file "{src}" as text)
  delay 1
  save workbook as wb filename (POSIX file "{pdf}" as text) file format PDF file format
  close wb saving no
end tell
end timeout'''
        r = subprocess.run(['osascript', '-e', script], capture_output=True, text=True)
        if not os.path.exists(pdf):
            raise SystemExit(f'{name}: export selhal {r.stderr}')
        import time
        for _ in range(20):                                         # Excel může dopisovat soubor ještě po návratu
            try:
                doc = fitz.open(pdf)
                break
            except Exception:
                time.sleep(0.5)
        doc.set_metadata({'title': name[:-5], 'author': 'REO HEATING' if root == A else 'Ria Trade, s.r.o.',
                          'creator': '', 'producer': ''})
        doc.save(f[:-5] + '.pdf', garbage=3, deflate=True)
        os.remove(pdf)
        os.remove(src)


if __name__ == '__main__':
    for d in (A, B):
        shutil.rmtree(d, ignore_errors=True)
    rows = {'A': [], 'B': []}
    for src in sorted(glob.glob(os.path.join(SRC_FORMS, '*.xlsx'))):
        code = re.search(r'PW\d\d-DKXLRE\d-H8', os.path.basename(src)).group(0)
        for var in 'AB':
            rows[var].append((code, form(src, var, code), label(code, var)))
    os.makedirs(os.path.join(B, 'certifikaty-keymark'), exist_ok=True)
    for f in sorted(glob.glob(os.path.join(SRC_CERTS, 'Cert 041-K032-*.pdf'))):
        shutil.copyfile(f, os.path.join(B, 'certifikaty-keymark', 'HP-KEYMARK_' + os.path.basename(f)[5:].replace(' ', '')))
    os.makedirs(os.path.join(B, 'protokoly-erp'), exist_ok=True)       # zkušební protokoly ErP (TÜV SÜD) výrobce
    for f in glob.glob(os.path.join(SRC_LABELS, '*.pdf')):
        if 'label' not in os.path.basename(f).lower():
            name = re.sub(r'[^\w.-]+', '_', os.path.basename(f)).replace('Helios', 'H8')
            shutil.copyfile(f, os.path.join(B, 'protokoly-erp', name))
    for var in 'AB':
        summary(var, rows[var])
    if '--pdf' in sys.argv:
        export_pdf(A)
        export_pdf(B)
    for d in (A, B):
        print(os.path.basename(d), sum(len(fs) for _, _, fs in os.walk(d)), 'souborů')
