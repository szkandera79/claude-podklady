#!/usr/bin/env python3
"""Česká prezentace REO HEATING z originálu výrobce.

    python3 extract.py      # texty k překladu → chunks/
    (překlad → cs/chunk_*.json podle PREKLAD-ZADANI.md)
    python3 images.py       # upravené obrázky → media_out/
    python3 build.py        # → REO-R290-Smart_prezentace-CZ.pptx

Co build dělá: dosadí překlad (se zachováním formátu runů a zalomení), přejmenuje modely, přebarví modré
prvky šablony na zelenou REO HEATING, odstraní vložené OLE objekty (Excel/PowerPoint výrobce), znovu postaví
snímky 2 (typové označení) a 4 (přehled vlastností), vymění loga/obrázky a přepíše metadata.
"""
import copy
import glob
import io
import json
import os
import re
import zipfile

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from common import *

OUT = os.path.join(DIR, 'REO-R290-Smart_prezentace-CZ.pptx')
GREEN = RGBColor(0x2E, 0x6B, 0x33)
GREEN_HEX = '2E6B33'
BLUES = {'00A2E9', '00B0F0', '0090E7', '0F9FE1'}       # modré barvy výrobce → zelená REO HEATING
INK = RGBColor(0x1F, 0x2A, 0x22)
GREY = RGBColor(0x5C, 0x6B, 0x60)
TINT = RGBColor(0xE7, 0xF1, 0xE6)

tr = {}
for f in sorted(glob.glob(os.path.join(DIR, 'cs', 'chunk_*.json'))):
    tr.update(json.load(open(f)))
seg = json.load(open(os.path.join(DIR, 'segments.json')))
tr.update({'1_1': 'Aktualizace: 20251212',
           '1_0': '<f1>Tepelné čerpadlo</f1><f2> </f2><f3>REO R290 Smart</f3><br><f4>Konstrukce a klíčové funkce</f4><br><f5>6/8/10/12/14/16 kW</f5>'})       # delší čeština se zalamovala pod okraj snímku

prs = Presentation(SRC)
A = lambda tag: qn('a:' + tag)
TOKEN = re.compile(r'<f(\d+)>(.*?)</f\1>|<br>|((?:(?!<br>|<f\d+>).)+)', re.S)


# ---------------------------------------------------------------- překlad
def write_back(p, cs):
    """Nahradí obsah odstavce českým textem; formát bere z původních runů (podle značek <fN>)."""
    g = groups(p)
    tg = [rs for k, rs in g if k != 'br']
    if not tg or any(r.tag == A('fld') for _, rs in g for r in rs):
        return False
    br_tpl = next((rs[0] for k, rs in g if k == 'br'), None)
    for k, rs in g:
        for r in rs:
            p.remove(r)
    anchor = p.find(A('endParaRPr'))
    new = []
    for m in TOKEN.finditer(cs):
        if m.group(0) == '<br>':
            new.append(copy.deepcopy(br_tpl) if br_tpl is not None else etree.SubElement(etree.Element('x'), A('br')))
            continue
        idx, txt = (int(m.group(1)) - 1, m.group(2)) if m.group(1) else (0, m.group(3))
        if not txt:
            continue
        r = copy.deepcopy(tg[min(idx, len(tg) - 1)][0])
        r.find(A('t')).text = txt
        rpr = r.find(A('rPr'))
        if rpr is not None:
            rpr.set('lang', 'cs-CZ')
            rpr.attrib.pop('altLang', None)
        new.append(r)
    for el in new:
        if anchor is not None:
            anchor.addprevious(el)
        else:
            p.append(el)
    return True


def visible_len(t):
    return len(re.sub(r'<[^>]+>', '', t))


applied, ratio = 0, {}
for sno, slide in enumerate(prs.slides, 1):
    if sno in REBUILT:
        continue
    for n, p, body in iter_paragraphs(slide):
        key = f'{sno}_{n}'
        if key in seg:
            cs = tr.get(key, seg[key])
            en = to_tagged(p)
            if write_back(p, cs):
                applied += 1
                r = ratio.setdefault(id(body), [body, 0, 0])
                r[1] += visible_len(en)
                r[2] += visible_len(cs)
        else:                                        # neprekládané odstavce: jen přejmenování modelů
            for rr in runs(p):
                t = rr.find(A('t'))
                if t is not None and t.text:
                    t.text = rename_models(t.text)

# ---------------------------------------------------------------- přizpůsobení délce (čeština je delší)
shrunk = 0
for body, en, cs in ratio.values():
    if en < 12 or cs <= en * 1.06:
        continue
    f = max(0.72, (en / cs) ** 0.8)
    if body.tag == A('txBody'):                       # buňka tabulky → zmenšit explicitní velikosti
        for rpr in body.iter(A('rPr'), A('endParaRPr')):
            if rpr.get('sz'):
                rpr.set('sz', str(max(700, int(int(rpr.get('sz')) * f / 50) * 50)))
    else:                                             # textové pole → automatické zmenšení písma
        bp = body.find(A('bodyPr'))
        for tag in ('spAutoFit', 'noAutofit', 'normAutofit'):
            for el in bp.findall(A(tag)):
                bp.remove(el)
        na = etree.SubElement(bp, A('normAutofit'))
        na.set('fontScale', str(int(f * 100000 / 2500) * 2500))
        if f < 0.9:
            na.set('lnSpcReduction', '10000')
    shrunk += 1


# ---------------------------------------------------------------- pomocné funkce pro nové tvary
def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def drop_unused_rels(slide):
    xml = etree.tostring(slide._element).decode()
    for rid, rel in list(slide.part.rels.items()):
        if rel.reltype.endswith(('/image', '/oleObject', '/package', '/vmlDrawing')) and f'"{rid}"' not in xml:
            slide.part.drop_rel(rid)


def text(slide, x, y, w, h, s, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, 'margin_' + side, Inches(0.04))
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        parts = line if isinstance(line, list) else [(line, bold)]
        for t, b in parts:
            r = para.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.name = 'Arial'
            r.font.color.rgb = color
    return tb


def box(slide, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    return s


def connector(slide, x1, y1, x2, y2, color=GREEN, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def set_title(slide, s):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() and sh.top < Inches(0.5):
            p = sh.text_frame.paragraphs[0]
            p.runs[0].text = s
            for r in p.runs[1:]:
                r.text = ''
            return sh


def media_blob(name):
    return io.BytesIO(zipfile.ZipFile(SRC).read('ppt/media/' + name))


# ---------------------------------------------------------------- snímek 2: typové označení
def unique_ids(slide, start):
    """Nové tvary dostanou ID jedinečná v celé prezentaci (náhled macOS jinak zaměňuje tvary mezi snímky)."""
    n = start
    for el in slide._element.iter(qn('p:cNvPr')):
        if el not in ORIG_IDS[slide.slide_id]:
            el.set('id', str(n))
            n += 1


s2 = prs.slides[1]
s4_ = prs.slides[3]
ORIG_IDS = {sl.slide_id: list(sl._element.iter(qn('p:cNvPr'))) for sl in (s2, s4_)}
set_title(s2, 'Typové označení')
for sh in list(s2.shapes):
    if sh.shape_type in (6, 13):                      # skupina (rozklad kódu výrobce) a čínský diagram řady
        remove_shape(sh)
tbl = next(sh for sh in s2.shapes if sh.has_table).table
for i, kw in enumerate(['06', '08', '10', '12', '14', '16']):
    for j, ph in enumerate(['1F', '3F']):
        cell = tbl.cell(i, j)
        para = cell.text_frame.paragraphs[0]
        val = '–' if (kw == '06' and ph == '3F') else f'RH{kw} {ph}'
        if para.runs:
            para.runs[0].text = val
            for r in para.runs[1:]:
                r.text = ''
        else:
            para.add_run().text = val
text(s2, 0.58, 3.08, 3.9, 0.3, '1F = jednofázové 230 V · 3F = třífázové 400 V', 11, color=GREY)

# rozklad označení (vlevo dole)
text(s2, 0.55, 3.5, 5.6, 0.35, 'Jak číst označení modelu', 15, True, GREEN)
code = [('RH', 0.8), ('08', 2.15), ('1F', 3.5)]
for t, x in code:
    box(s2, x, 3.92, 1.05, 0.68, GREEN)
    text(s2, x, 3.92, 1.05, 0.68, t, 26, True, RGBColor(255, 255, 255), PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
legend = [('RH', 'REO HEATING – tepelné čerpadlo vzduch–voda'),
          ('08', 'jmenovitý topný výkon v kW (06–16)'),
          ('1F / 3F', 'napájení: 1F = 230 V, 3F = 400 V')]
for i, (k, v) in enumerate(legend):
    y = 4.8 + i * 0.4
    connector(s2, 0.62, y + 0.17, 0.72, y + 0.17)
    text(s2, 0.78, y, 5.4, 0.38, [[(k + ': ', True), (v, False)]], 13)
text(s2, 0.58, 6.12, 5.6, 0.9, ['Typ: nízkoteplotní tepelné čerpadlo vzduch–voda',
                               'Funkce: vytápění, chlazení, ohřev TUV · Chladivo: R290'], 12, color=GREY)

# produktová řada (vpravo) – místo čínského diagramu
X0, Y0 = 6.3, 1.25
box(s2, X0, Y0 + 1.75, 1.9, 0.95, GREEN)
text(s2, X0, Y0 + 1.75, 1.9, 0.95, ['REO R290', 'Smart'], 15, True, RGBColor(255, 255, 255),
     PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
branches = [
    (0.55, 'Rezidenční', '6–16 kW', 'RH06 – RH16', 'S vestavěným hydraulickým modulem · SCOP až 5,3'),
    (2.35, None, None, 'RH06 – RH16', 'Verze s vestavěným oběhovým čerpadlem · SCOP až 5,3'),
    (4.15, 'Komerční', '20 / 30 kW', 'RHX20 Pro / RHX30 Pro', 'Lehké komerční použití · A+++ · nízká hlučnost'),
]
connector(s2, X0 + 1.9, Y0 + 2.22, X0 + 2.2, Y0 + 2.22)
connector(s2, X0 + 2.2, Y0 + 0.95, X0 + 2.2, Y0 + 4.55)
for dy, seg_name, rng, model, desc in branches:
    yc = Y0 + dy + 0.4
    connector(s2, X0 + 2.2, yc, X0 + 2.45, yc)
    if seg_name:
        text(s2, X0 + 2.45, Y0 + dy, 1.35, 0.8, [[(seg_name, True)], [(rng, False)]], 13, color=GREEN,
             anchor=MSO_ANCHOR.MIDDLE)
    box(s2, X0 + 3.85, Y0 + dy, 2.75, 0.8, TINT)
    text(s2, X0 + 3.95, Y0 + dy + 0.02, 2.6, 0.78, [[(model, True)], [(desc, False)]], 11,
         anchor=MSO_ANCHOR.MIDDLE)
drop_unused_rels(s2)
unique_ids(s2, 9001)

# ---------------------------------------------------------------- snímek 4: přehled vlastností
s4 = prs.slides[3]
set_title(s4, 'Přehled vlastností')
for sh in list(s4.shapes):
    if sh.shape_type in (7, 13):                      # srovnávací obrázek s řadami výrobce + vložený Excel
        remove_shape(sh)
    elif sh.has_text_frame and 'hydraulic' in sh.text_frame.text:
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = 'Jednotky mají vestavěný hydraulický modul'
        for r in p.runs[1:]:
            r.text = ''
rows = [
    ('Konstrukce', 'Jeden ventilátor: 6 / 8 / 10 / 12 / 14 / 16 kW'),
    ('Konfigurace', 'Kompresor MITSUBISHI, oběhové čerpadlo SHIMGE, hydraulický modul '
                    '(elektrický ohřev 3 kW, expanzní nádoba 6 l, pojistný a odvzdušňovací ventil)'),
    ('Sezónní účinnost', 'SCOP až 5,30'),
    ('Teplota výstupní vody', 'až 78 °C'),
    ('Přímé napájení z FV', 'deska MPPT – volitelně'),
    ('Snímač úniku chladiva', 'standardně'),
    ('Odvzdušňovací ventil', 'vestavěný'),
    ('Ovladač', '7" barevný dotykový displej'),
]
gt = s4.shapes.add_table(len(rows), 2, Inches(0.55), Inches(1.4), Inches(7.6), Inches(5.4)).table
gt.columns[0].width = Inches(2.45)
gt.columns[1].width = Inches(5.15)
for i, (k, v) in enumerate(rows):
    for j, t in enumerate((k, v)):
        c = gt.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = TINT if j == 0 else RGBColor(255, 255, 255)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame
        tf.paragraphs[0].text = ''
        r = tf.paragraphs[0].add_run()
        r.text = t
        r.font.size = Pt(14)
        r.font.bold = j == 0
        r.font.name = 'Arial'
        r.font.color.rgb = GREEN if j == 0 else INK
tblpr = gt._tbl.tblPr
tblpr.set('firstRow', '0')
tblpr.set('bandRow', '0')
s4.shapes.add_picture(media_blob('image5.png'), Inches(8.55), Inches(1.35), width=Inches(4.2))
s4.shapes.add_picture(media_blob('image27.png'), Inches(8.75), Inches(5.05), width=Inches(3.8))
drop_unused_rels(s4)
unique_ids(s4, 9501)

# ---------------------------------------------------------------- vložené OLE objekty výrobce (snímky 30, 39)
for sno, label in ((30, 'User parameters'), (39, 'IoT')):
    sl = prs.slides[sno - 1]
    for sh in list(sl.shapes):
        if sh.shape_type == 7:
            remove_shape(sh)
        elif sh.has_text_frame and sh.left > Inches(10) and label.split()[0] in sh.text_frame.text + seg.get('', ''):
            pass
    # popisky ikon OLE („Seznam uživatelských parametrů“, „Odkaz IoT“) – už nemají na co ukazovat
    for sh in list(sl.shapes):
        if sh.has_text_frame and sh.left >= Inches(10.1) and sh.top < Inches(4.8) and sh.width < Inches(2.5):
            remove_shape(sh)
    drop_unused_rels(sl)

# ---------------------------------------------------------------- snímek 9: obrázek izolace bez čínských popisků
for sh in prs.slides[8].shapes:
    if sh.shape_type == 13 and sh.image.filename and abs(sh.width - Inches(2.29)) < Inches(0.05):
        sh.height = int(sh.width * 515 / 1388)

# ---------------------------------------------------------------- doladění podle exportu z PowerPointu
# tabulky technických údajů (snímky 5, 6): čeština se v úzkých sloupcích víc zalamuje → menší písmo, bez horního okraje
SZ_MAP = {'800': '680', '750': '640', '700': '600'}
for sno in (5, 6):
    tbl = next(sh for sh in prs.slides[sno - 1].shapes if sh.has_table)
    for el in tbl._element.iter(A('rPr'), A('endParaRPr')):
        if el.get('sz') in SZ_MAP:
            el.set('sz', SZ_MAP[el.get('sz')])
    for tcpr in tbl._element.iter(A('tcPr')):
        tcpr.set('marT', '0')
        tcpr.set('marB', '0')
# snímek 25: delší výčet jazyků → menší písmo textového pole
for sh in prs.slides[24].shapes:
    if sh.has_text_frame and 'jazyk' in sh.text_frame.text:
        for el in sh._element.iter(A('rPr'), A('endParaRPr')):
            if el.get('sz') == '1600':
                el.set('sz', '1400')
# snímek 3: napájení – kratší mezera mezi hodnotami (jinak se 3f hodnota ořízne)
for sh in prs.slides[2].shapes:
    if sh.has_table:
        for row in sh.table.rows:
            for c in row.cells:
                for t in c._tc.iter(A('t')):
                    if t.text and '380V/3Ph' in t.text:
                        t.text = re.sub(r'\s{4,}', '          ', t.text)

# snímek 3: tabulka je česky nižší → řádek napájení vyjel pod fotky jednotek; fotky končící níž než 4,72" zmenšit
LIMIT = Inches(4.72)
for sh in prs.slides[2].shapes:
    if sh.shape_type == 13 and sh.top + sh.height > LIMIT:
        f = (LIMIT - sh.top) / sh.height
        cx = sh.left + sh.width / 2
        sh.width, sh.height = int(sh.width * f), int(sh.height * f)
        sh.left = int(cx - sh.width / 2)

# ---------------------------------------------------------------- barvy, písma, alt texty
def recolor(root, sno=0):
    for c in root.iter(A('srgbClr')):
        if c.get('val', '').upper() in BLUES:
            c.set('val', GREEN_HEX)
    for lat in root.iter(A('latin')):
        if lat.get('typeface') in ('微软雅黑', '等线', '等线 Light', '宋体'):
            lat.set('typeface', 'Arial')
    for nv in root.iter(qn('p:cNvPr')):            # čínské alt texty → jedinečný český popis
        if nv.getparent().tag == qn('p:nvPicPr') or nv.get('descr'):
            nv.set('descr', f'Obrázek – snímek {sno}, prvek {nv.get("id")}' if sno else 'Logo REO HEATING')


for i, sl in enumerate(prs.slides, 1):
    recolor(sl._element, i)
for layout in prs.slide_layouts:
    recolor(layout._element)
recolor(prs.slide_master._element)
theme_part = next(r.target_part for r in prs.slide_master.part.rels.values() if r.reltype.endswith('/theme'))
theme = etree.fromstring(theme_part.blob)
for fs in theme.iter(A('majorFont'), A('minorFont')):
    fs.find(A('latin')).set('typeface', 'Arial')
for nm in theme.iter(A('theme')):
    nm.set('name', 'REO HEATING')
theme_part._blob = etree.tostring(theme, xml_declaration=True, encoding='UTF-8', standalone=True)

cp = prs.core_properties
cp.title = 'REO R290 Smart – tepelná čerpadla RH06–RH16'
cp.author = cp.last_modified_by = 'REO HEATING'
cp.subject = 'REO R290 Smart – konstrukce a hlavní funkce invertorových tepelných čerpadel vzduch–voda'
cp.keywords = 'REO R290 Smart, REO HEATING, RH06, RH08, RH10, RH12, RH14, RH16, R290, tepelné čerpadlo'
cp.category = cp.comments = ''
cp.revision = 1
tmp = OUT + '.tmp'
prs.save(tmp)

# ---------------------------------------------------------------- balíček: vyměnit obrázky, přepsat app.xml, titulní pozadí
# python-pptx při uložení přečísluje obrázky → dosazovat podle obsahu (otisk originálního souboru), ne podle jména
import hashlib
zsrc = zipfile.ZipFile(SRC)
MEDIA = {}
for f in glob.glob(os.path.join(DIR, 'media_out', '*')):
    orig = zsrc.read('ppt/media/' + os.path.basename(f))
    MEDIA[hashlib.sha1(orig).hexdigest()] = f
APP = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
       '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
       'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
       f'<Application>Microsoft Office PowerPoint</Application><Slides>{len(prs.slides)}</Slides>'
       '<Company>REO HEATING</Company></Properties>')
zin = zipfile.ZipFile(tmp)
replaced = []
# souvislé číslování ppt/media/imageN (náhled macOS přiřazuje obrázky podle pořadí; mezery po odstraněných
# obrázcích výrobce by v něm obrázky posunuly)
media = sorted((n for n in zin.namelist() if re.match(r'ppt/media/image\d+\.', n)),
               key=lambda n: int(re.search(r'image(\d+)', n).group(1)))
RENAME = {old: f'ppt/media/image{i}{os.path.splitext(old)[1]}' for i, old in enumerate(media, 1)}


def fix_rels(xml):
    return re.sub(r'Target="\.\./media/(image\d+\.\w+)"',
                  lambda m: 'Target="../media/' + os.path.basename(RENAME.get('ppt/media/' + m.group(1),
                                                                               'ppt/media/' + m.group(1))) + '"', xml)
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        name = os.path.basename(item.filename)
        h = hashlib.sha1(data).hexdigest()
        if item.filename.startswith('ppt/media/') and h in MEDIA:
            new = MEDIA[h]
            if os.path.splitext(new)[1] != os.path.splitext(name)[1]:
                raise SystemExit(f'nesouhlasí typ obrázku: {name} ← {new}')
            data = open(new, 'rb').read()
            replaced.append(os.path.basename(new))
        elif item.filename == 'docProps/app.xml':
            data = APP.encode()
        elif item.filename.endswith('.rels'):
            data = fix_rels(data.decode('utf8')).encode('utf8')
        zout.writestr(RENAME.get(item.filename, item.filename), data)
os.remove(tmp)
print('vyměněné obrázky:', sorted(replaced))
print(f'odstavců přeloženo: {applied} | zmenšeno textových polí/buněk: {shrunk} | uloženo: {OUT}')
