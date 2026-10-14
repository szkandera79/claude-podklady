#!/usr/bin/env python3
"""Česká prezentace REO R290 Smart RHX20 Pro / RHX30 Pro (REO HEATING) z originálu výrobce.

    python3 extract.py      # texty k překladu → chunks/
    (překlad → cs/chunk_*.json podle PREKLAD-ZADANI.md)
    python3 images.py       # upravené obrázky → media_out/
    python3 build.py        # → REO-R290-Smart_RHX20-RHX30-Pro_prezentace-CZ.pptx

Co build dělá: dosadí překlad (se zachováním formátu runů a zalomení; reused.json = shodné texty přeložené
u prezentace RH06–RH16, cs/ = nové), přejmenuje modely, přebarví modré prvky šablony na zelenou REO HEATING,
vložené OLE objekty nahradí obrázkem (schémata ZWCAD) nebo odstraní (Excel), vymění loga/obrázky
a přepíše metadata.
"""
import copy
import glob
import io
import json
import os
import re
import zipfile

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from common import *

OUT = os.path.join(DIR, 'REO-R290-Smart_RHX20-RHX30-Pro_prezentace-CZ.pptx')
GREEN = RGBColor(0x2E, 0x6B, 0x33)
GREEN_HEX = '2E6B33'
BLUES = {'00A2E9', '00B0F0', '0090E7', '0F9FE1', '1067A9'}       # modré barvy výrobce → zelená REO HEATING
INK = RGBColor(0x1F, 0x2A, 0x22)
GREY = RGBColor(0x5C, 0x6B, 0x60)
TINT = RGBColor(0xE7, 0xF1, 0xE6)

tr = json.load(open(os.path.join(DIR, 'reused.json')))
for f in sorted(glob.glob(os.path.join(DIR, 'cs', 'chunk_*.json'))):
    tr.update(json.load(open(f)))
seg = json.load(open(os.path.join(DIR, 'segments.json')))
tr.update({'1_0': '<f1>Tepelné čerpadlo REO R290 Smart</f1><br><f2>Konstrukce a klíčové funkce</f2><br>'
                  '<f3>RHX20 Pro / RHX30 Pro · lehké komerční použití</f3>',
           '24_1': 'Hlavní obrazovka → User parameters      Kompletní seznam uživatelských parametrů je v instalačním návodu. '})

prs = Presentation(SRC)
A = lambda tag: qn('a:' + tag)
TOKEN = re.compile(r'<f(\d+)>(.*?)</f\1>|<br>|((?:(?!<br>|<f\d+>).)+)', re.S)


# ---------------------------------------------------------------- překlad
def write_back(p, cs):
    """Nahradí obsah odstavce českým textem; formát bere z původních runů (podle značek <fN>)."""
    g = groups(p)
    tg = [rs for k, rs in g if k != 'br']
    if not tg or any(r.tag == A('fld') for _, rs in g for r in rs):
        return False
    br_tpl = next((rs[0] for k, rs in g if k == 'br'), None)
    for k, rs in g:
        for r in rs:
            p.remove(r)
    anchor = p.find(A('endParaRPr'))
    new = []
    for m in TOKEN.finditer(cs):
        if m.group(0) == '<br>':
            new.append(copy.deepcopy(br_tpl) if br_tpl is not None else etree.SubElement(etree.Element('x'), A('br')))
            continue
        idx, txt = (int(m.group(1)) - 1, m.group(2)) if m.group(1) else (0, m.group(3))
        if not txt:
            continue
        r = copy.deepcopy(tg[min(idx, len(tg) - 1)][0])
        r.find(A('t')).text = txt
        rpr = r.find(A('rPr'))
        if rpr is not None:
            rpr.set('lang', 'cs-CZ')
            rpr.attrib.pop('altLang', None)
        new.append(r)
    for el in new:
        if anchor is not None:
            anchor.addprevious(el)
        else:
            p.append(el)
    return True


def visible_len(t):
    return len(re.sub(r'<[^>]+>', '', t))


applied, ratio = 0, {}
for sno, slide in enumerate(prs.slides, 1):
    if sno in REBUILT:
        continue
    for n, p, body in iter_paragraphs(slide):
        key = f'{sno}_{n}'
        if key in seg:
            cs = tr.get(key, seg[key])
            en = to_tagged(p)
            if write_back(p, cs):
                applied += 1
                r = ratio.setdefault(id(body), [body, 0, 0])
                r[1] += visible_len(en)
                r[2] += visible_len(cs)
        else:                                        # neprekládané odstavce: jen přejmenování modelů
            for rr in runs(p):
                t = rr.find(A('t'))
                if t is not None and t.text:
                    t.text = rename_models(t.text)

# ---------------------------------------------------------------- přizpůsobení délce (čeština je delší)
shrunk = 0
for body, en, cs in ratio.values():
    if en < 12 or cs <= en * 1.06:
        continue
    f = max(0.72, (en / cs) ** 0.8)
    if body.tag == A('txBody'):                       # buňka tabulky → zmenšit explicitní velikosti
        for rpr in body.iter(A('rPr'), A('endParaRPr')):
            if rpr.get('sz'):
                rpr.set('sz', str(max(700, int(int(rpr.get('sz')) * f / 50) * 50)))
    else:                                             # textové pole → automatické zmenšení písma
        bp = body.find(A('bodyPr'))
        for tag in ('spAutoFit', 'noAutofit', 'normAutofit'):
            for el in bp.findall(A(tag)):
                bp.remove(el)
        na = etree.SubElement(bp, A('normAutofit'))
        na.set('fontScale', str(int(f * 100000 / 2500) * 2500))
        if f < 0.9:
            na.set('lnSpcReduction', '10000')
    shrunk += 1



def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def drop_unused_rels(slide):
    xml = etree.tostring(slide._element).decode()
    for rid, rel in list(slide.part.rels.items()):
        if rel.reltype.endswith(('/image', '/oleObject', '/package', '/vmlDrawing')) and f'"{rid}"' not in xml:
            slide.part.drop_rel(rid)


P = lambda tag: qn('p:' + tag)

# ---------------------------------------------------------------- vložené OLE objekty
# snímky 34, 35: schémata zapojení (výkres ZWCAD) → statický obrázek z náhledu OLE (stejná poloha a velikost)
for sno in (34, 35):
    sl = prs.slides[sno - 1]
    for gf in list(sl._element.iter(P('graphicFrame'))):
        pic = next(gf.iter(P('pic')), None)
        if pic is None:
            continue
        pic = copy.deepcopy(pic)
        nv = pic.find(P('nvPicPr')).find(P('cNvPr'))
        old = gf.find(P('nvGraphicFramePr')).find(P('cNvPr'))
        nv.set('id', old.get('id'))
        nv.set('name', 'Schéma zapojení')
        gf.addprevious(pic)
        gf.getparent().remove(gf)
# snímek 24: vložený Excel výrobce se seznamem parametrů + jeho popisek → pryč (odkaz na návod je v textu)
s24 = prs.slides[23]
for sh in list(s24.shapes):
    if sh.shape_type == 7 or (sh.has_text_frame and sh.left > Inches(10.5) and sh.width < Inches(2.5)):
        remove_shape(sh)
for sno in (24, 34, 35):
    sl = prs.slides[sno - 1]
    for ld in list(sl._element.iter(P('legacyDrawing'))):
        ld.getparent().remove(ld)
    drop_unused_rels(sl)

# ---------------------------------------------------------------- doladění podle exportu z PowerPointu
# podnadpisy „Vnitřní řídicí deska …“ (snímky 34–38) mají dost místa → bez automatického zmenšení
for sno in range(34, 39):
    for sh in prs.slides[sno - 1].shapes:
        if sh.has_text_frame and 'řídicí deska' in sh.text_frame.text.lower():
            bp = sh.text_frame._txBody.find(A('bodyPr'))
            for el in bp.findall(A('normAutofit')):
                bp.remove(el)
# snímek 13: popisek „Řídicí deska“ je delší než „Mainboard“ → posunout vlevo, aby nezasahoval do konektoru
for sh in prs.slides[12].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == 'Řídicí deska':
        sh.left -= Inches(0.5)
        sh.width += Inches(0.5)
# snímek 3: česká tabulka je nižší → fotky jednotek přesahovaly do řádku „Napájení“ → zmenšit (spodní okraj 4,45")
LIMIT = Inches(4.45)
for sh in prs.slides[2].shapes:
    if sh.shape_type == 13 and sh.top + sh.height > LIMIT:
        f = (LIMIT - sh.top) / sh.height
        cx = sh.left + sh.width / 2
        sh.width, sh.height = int(sh.width * f), int(sh.height * f)
        sh.left = int(cx - sh.width / 2)
# snímek 44: „Děkujeme!“ se v úzkém poli zalamovalo
for sh in prs.slides[43].shapes:
    if sh.has_text_frame and 'Děkujeme' in sh.text_frame.text:
        cx = sh.left + sh.width // 2
        sh.width = Inches(9)
        sh.left = int(cx - sh.width / 2)
        sh.text_frame.word_wrap = False

# ---------------------------------------------------------------- snímky 1 a 44: bílé pozadí, zelené písmo
# skryje grafiku layoutu (zelená mapa + bílé logo), nastaví bílé pozadí a vloží barevné logo REO HEATING
def white_slide(s1):
    s1._element.set('showMasterSp', '0')
    s1.background.fill.solid()
    s1.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    logo = os.path.join(DIR, 'media_out', 'image3.png')
    lw = Emu(1588444)
    pic = s1.shapes.add_picture(logo, Emu(10051576), Emu(344557), width=lw)
    pic._element.getparent().remove(pic._element)
    s1.shapes._spTree.insert(2, pic._element)                  # za kořen stromu, pod texty
    for sh in s1.shapes:
        if not sh.has_text_frame:
            continue
        for rpr in sh._element.iter(A('rPr'), A('endParaRPr')):
            for el in rpr.findall(A('solidFill')) + rpr.findall(A('effectLst')):
                rpr.remove(el)
            fill = etree.SubElement(rpr, A('solidFill'))
            etree.SubElement(fill, A('srgbClr')).set('val', GREEN_HEX)
            rpr.insert(0, fill)                                # solidFill musí být v rPr před latin/ea/cs
    for el in s1._element.iter(qn('p:cNvPr')):
        if el.get('descr') is None and el.getparent().tag == qn('p:nvPicPr'):
            el.set('descr', 'Logo REO HEATING')


for sl in (prs.slides[0], prs.slides[43]):          # titulní a závěrečný snímek
    white_slide(sl)

# ---------------------------------------------------------------- snímek 4: hodnoty v ikonách infografiky
# max. teplota vody: výrobce v ikoně uvádí 78 °C, v tabulce 25~75 °C → jednotně 75 (pravidlo REO R290 Smart); řeší images.py
for grp in prs.slides[3]._element.iter(P('grpSp')):
    if next(grp.iter(P('pic')), None) is None or grp.getparent().tag != P('spTree'):
        continue
    for sp in list(grp.iter(P('sp'))):                      # překryvy „78“/„5.0“: hodnoty jsou nově přímo v obrázku
        sp.getparent().remove(sp)

def text(slide, x, y, w, h, s, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, 'margin_' + side, Inches(0.04))
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        parts = line if isinstance(line, list) else [(line, bold)]
        for t, b in parts:
            r = para.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.name = 'Arial'
            r.font.color.rgb = color
    return tb


# ---------------------------------------------------------------- nový snímek 5: výkon při −15 °C (výkonové mapy výrobce)
# vkládá se až po všech úpravách podle čísel snímků; fix_icons.py proto počítá s posunem o +1 od snímku 5
MAPS = json.load(open(os.path.join(DIR, '..', 'data', 'vykonove-mapy.json')))
ns = prs.slides.add_slide(prs.slides[3].slide_layout)
for ph in list(ns.placeholders):
    if ph.placeholder_format.type != 1:                     # ponechat jen nadpis
        remove_shape(ph)
ns.shapes.title.text = 'Výkon při −15 °C'
for r in ns.shapes.title.text_frame.paragraphs[0].runs:
    r.font.name = 'Times New Roman'
t4 = next(sh for sh in prs.slides[3].shapes if sh.has_text_frame and sh.top < Inches(0.5))
tt = ns.shapes.title
tt.left, tt.top, tt.width, tt.height = t4.left, t4.top, Inches(10), t4.height
fmt = lambda v, d=1: f'{v:.{d}f}'.replace('.', ',')
# jednotný přehled: max. topný výkon a COP při A−15 pro W35 (podlahovka) a W55 (radiátory)
HDR1 = ['Model', 'A−15 / W35', '', 'A−15 / W55', '']
HDR2 = ['', 'Topný výkon', 'COP', 'Topný výkon', 'COP']
rows = []
for m in ('RHX20 Pro', 'RHX30 Pro'):
    h = MAPS[m]['heating']
    i = h['amb'].index(-15.0)
    r = [m]
    for w in ('35', '55'):
        mx = h[w]['max']
        r += [f'{fmt(mx["cap"][i])} kW', fmt(mx['cop'][i], 2)]
    rows.append(r)
W = [2.6, 2.3, 1.5, 2.3, 1.5]
gt = ns.shapes.add_table(len(rows) + 2, len(W), Inches(1.55), Inches(1.6), Inches(sum(W)), Inches(2.6)).table
for j, w in enumerate(W):
    gt.columns[j].width = Inches(w)
for i, row in enumerate([HDR1, HDR2] + rows):
    for j, t in enumerate(row):
        c = gt.cell(i, j)
        c.fill.solid()
        hdr = i < 2
        c.fill.fore_color.rgb = GREEN if hdr else (TINT if j == 0 else RGBColor(255, 255, 255))
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = c.text_frame.paragraphs[0]
        para.text = ''
        para.alignment = PP_ALIGN.CENTER
        r = para.add_run()
        r.text = t
        r.font.size = Pt(16 if hdr else 22 if j in (1, 3) else 18)
        r.font.name = 'Arial'
        r.font.bold = hdr or j in (0, 1, 3)
        r.font.color.rgb = RGBColor(255, 255, 255) if hdr else (GREEN if j in (1, 3) else INK)
gt.cell(0, 0).merge(gt.cell(1, 0))
gt.cell(0, 1).merge(gt.cell(0, 2))
gt.cell(0, 3).merge(gt.cell(0, 4))
tblpr = gt._tbl.tblPr
tblpr.set('firstRow', '0')
tblpr.set('bandRow', '0')
text(ns, 1.55, 4.4, 10.3, 0.75, [[('Číslo v označení (20 / 30) je jmenovitý výkon při A7/W35, ne výkon v mrazu. ', True),
                                    ('Pro dimenzování podle tepelné ztráty používejte hodnoty z tabulky.', False)]],
     15, color=GREEN)
text(ns, 1.55, 5.3, 10.3, 1.2, [
    'A−15 = venkovní vzduch −15 °C · W35 = výstupní voda 35 °C (podlahové vytápění) · '
    'W55 = výstupní voda 55 °C (radiátory). Uveden maximální topný výkon jednotky.',
    'Zdroj: výkonové mapy výrobce. Hodnoty jsou orientační; skutečný výkon závisí na podmínkách instalace.'],
    12, color=GREY)
lst = prs.slides._sldIdLst                                  # přesunout za snímek 4 (technické údaje)
el = lst[-1]
lst.remove(el)
lst.insert(4, el)

# ---------------------------------------------------------------- barvy, písma, alt texty
def recolor(root, sno=0):
    for c in root.iter(A('srgbClr')):
        if c.get('val', '').upper() in BLUES:
            c.set('val', GREEN_HEX)
    for lat in root.iter(A('latin')):
        if lat.get('typeface') in ('微软雅黑', '等线', '等线 Light', '宋体'):
            lat.set('typeface', 'Arial')
    for nv in root.iter(qn('p:cNvPr')):            # čínské alt texty → jedinečný český popis
        if nv.getparent().tag == qn('p:nvPicPr') or nv.get('descr'):
            nv.set('descr', f'Obrázek – snímek {sno}, prvek {nv.get("id")}' if sno else 'Logo REO HEATING')


for i, sl in enumerate(prs.slides, 1):
    recolor(sl._element, i)
for layout in prs.slide_layouts:
    recolor(layout._element)
recolor(prs.slide_master._element)
theme_part = next(r.target_part for r in prs.slide_master.part.rels.values() if r.reltype.endswith('/theme'))
theme = etree.fromstring(theme_part.blob)
for fs in theme.iter(A('majorFont'), A('minorFont')):
    fs.find(A('latin')).set('typeface', 'Arial')
for nm in theme.iter(A('theme')):
    nm.set('name', 'REO HEATING')
theme_part._blob = etree.tostring(theme, xml_declaration=True, encoding='UTF-8', standalone=True)

cp = prs.core_properties
cp.title = 'REO R290 Smart – tepelná čerpadla RHX20 Pro / RHX30 Pro'
cp.author = cp.last_modified_by = 'REO HEATING'
cp.subject = 'REO R290 Smart RHX20 Pro / RHX30 Pro – konstrukce a hlavní funkce tepelných čerpadel vzduch–voda pro lehké komerční použití'
cp.keywords = 'REO R290 Smart, REO HEATING, RHX20 Pro, RHX30 Pro, R290, tepelné čerpadlo, kaskáda'
cp.category = cp.comments = ''
cp.revision = 1
tmp = OUT + '.tmp'
prs.save(tmp)

# ---------------------------------------------------------------- balíček: vyměnit obrázky, přepsat app.xml, titulní pozadí
# python-pptx při uložení přečísluje obrázky → dosazovat podle obsahu (otisk originálního souboru), ne podle jména
import hashlib
zsrc = zipfile.ZipFile(SRC)
MEDIA = {}
for f in glob.glob(os.path.join(DIR, 'media_out', '*')):
    orig = zsrc.read('ppt/media/' + os.path.basename(f))
    MEDIA[hashlib.sha1(orig).hexdigest()] = f
APP = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
       '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
       'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
       f'<Application>Microsoft Office PowerPoint</Application><Slides>{len(prs.slides)}</Slides>'
       '<Company>REO HEATING</Company></Properties>')
zin = zipfile.ZipFile(tmp)
replaced = []
# souvislé číslování ppt/media/imageN (náhled macOS přiřazuje obrázky podle pořadí; mezery po odstraněných
# obrázcích výrobce by v něm obrázky posunuly)
media = sorted((n for n in zin.namelist() if re.match(r'ppt/media/image\d+\.', n)),
               key=lambda n: int(re.search(r'image(\d+)', n).group(1)))
RENAME = {old: f'ppt/media/image{i}{os.path.splitext(old)[1]}' for i, old in enumerate(media, 1)}


def fix_rels(xml):
    return re.sub(r'Target="\.\./media/(image\d+\.\w+)"',
                  lambda m: 'Target="../media/' + os.path.basename(RENAME.get('ppt/media/' + m.group(1),
                                                                               'ppt/media/' + m.group(1))) + '"', xml)
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        name = os.path.basename(item.filename)
        h = hashlib.sha1(data).hexdigest()
        if item.filename.startswith('ppt/media/') and h in MEDIA:
            new = MEDIA[h]
            if os.path.splitext(new)[1] != os.path.splitext(name)[1]:
                raise SystemExit(f'nesouhlasí typ obrázku: {name} ← {new}')
            data = open(new, 'rb').read()
            replaced.append(os.path.basename(new))
        elif item.filename == 'docProps/app.xml':
            data = APP.encode()
        elif item.filename.endswith('.rels'):
            data = fix_rels(data.decode('utf8')).encode('utf8')
        zout.writestr(RENAME.get(item.filename, item.filename), data)
os.remove(tmp)
print('vyměněné obrázky:', sorted(replaced))
print(f'odstavců přeloženo: {applied} | zmenšeno textových polí/buněk: {shrunk} | uloženo: {OUT}')
